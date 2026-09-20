"""Settings dialog data tab: backup/restore and Excel import/export (SRP: data tab only)."""

from typing import Any, cast

from PyQt6.QtWidgets import (
    QFileDialog, QGroupBox, QLabel, QMessageBox, QPushButton, QVBoxLayout, QWidget,
)

import config
from ui.export_helpers import export_flights_to_excel


try:
    import openpyxl

    HAS_OPENPYXL = True
except ImportError:
    openpyxl = None
    HAS_OPENPYXL = False


def build_data_tab(dialog: Any) -> QWidget:
    widget = QWidget()
    layout = QVBoxLayout(widget)

    # 프로필 가져오기/내보내기 그룹
    grp_profile = QGroupBox("📦 설정 백업/복원")
    gp_layout = QVBoxLayout(grp_profile)

    btn_export_settings = QPushButton("💾 모든 설정 내보내기")
    btn_export_settings.setToolTip("프리셋, 프로필, 선호 시간대, 테마 등 모든 설정을 JSON 파일로 저장")
    btn_export_settings.clicked.connect(dialog._export_all_settings)

    btn_import_settings = QPushButton("📂 설정 가져오기")
    btn_import_settings.setToolTip("JSON 파일에서 설정을 불러와 현재 설정에 병합")
    btn_import_settings.clicked.connect(dialog._import_all_settings)

    gp_layout.addWidget(btn_export_settings)
    gp_layout.addWidget(btn_import_settings)
    layout.addWidget(grp_profile)

    # 엑셀 그룹
    grp_excel = QGroupBox("📊 엑셀 (Excel)")
    gl = QVBoxLayout(grp_excel)

    btn_import = QPushButton("📂 검색 조건 가져오기 (Import)")
    btn_import.clicked.connect(dialog._import_excel)

    label_info = QLabel("엑셀 파일 양식: Origin, Dest, DepDate(YYYYMMDD), RetDate, Adults")
    label_info.setStyleSheet("font-size: 11px; color: #aaa;")

    btn_export = QPushButton("💾 검색 결과 내보내기 (Export)")
    btn_export.clicked.connect(dialog._export_excel)

    gl.addWidget(btn_import)
    gl.addWidget(btn_export)
    gl.addWidget(label_info)

    layout.addWidget(grp_excel)
    layout.addStretch()
    return widget


def export_all_settings(dialog: Any) -> None:
    """모든 설정을 JSON 파일로 내보내기"""
    from datetime import datetime

    filename, _ = QFileDialog.getSaveFileName(
        dialog, "설정 내보내기",
        f"flight_bot_settings_{datetime.now().strftime('%Y%m%d')}.json",
        "JSON Files (*.json)"
    )
    if not filename:
        return

    if dialog.prefs.export_all_settings(filename):
        QMessageBox.information(dialog, "완료", f"설정이 저장되었습니다:\n{filename}")
    else:
        QMessageBox.critical(dialog, "오류", "설정 내보내기에 실패했습니다.")


def import_all_settings(dialog: Any) -> None:
    """JSON 파일에서 설정 가져오기"""
    filename, _ = QFileDialog.getOpenFileName(
        dialog, "설정 가져오기",
        "",
        "JSON Files (*.json)"
    )
    if not filename:
        return

    reply = QMessageBox.question(
        dialog, "설정 가져오기",
        "현재 설정에 가져온 설정이 병합됩니다.\n계속하시겠습니까?",
        QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
    )
    if reply != QMessageBox.StandardButton.Yes:
        return

    if dialog.prefs.import_settings(filename):
        QMessageBox.information(dialog, "완료", "설정을 가져왔습니다.\n변경 사항을 적용하려면 프로그램을 재시작하세요.")
        dialog._refresh_presets()
    else:
        QMessageBox.critical(dialog, "오류", "설정 가져오기에 실패했습니다.")


def import_excel(dialog: Any) -> None:
    if not HAS_OPENPYXL:
        QMessageBox.critical(dialog, "오류", "openpyxl 라이브러리가 설치되지 않았습니다.\npip install openpyxl")
        return

    fname, _ = QFileDialog.getOpenFileName(dialog, "엑셀 파일 열기", "", "Excel Files (*.xlsx)")
    if not fname:
        return

    try:
        if openpyxl is None:
            raise RuntimeError("openpyxl is unavailable")
        wb = openpyxl.load_workbook(fname)
        ws = wb.active
        if ws is None:
            raise RuntimeError("worksheet initialization failed")
        # Assume Row 2: Origin, Dest, DepDate, RetDate, Adults
        origin = ws['A2'].value
        dest = ws['B2'].value
        dep = str(ws['C2'].value)
        ret = str(ws['D2'].value) if ws['D2'].value else None
        adults = ws['E2'].value

        if origin and dest:
            params = {
                "origin": origin,
                "dest": dest,
                "dep": dep,
                "ret": ret,
                "adults": int(adults) if adults else 1,
                "cabin_class": "ECONOMY",
                "is_domestic": config.infer_is_domestic_route(origin, dest),
            }
            dialog.prefs.save_profile("엑셀 가져옴", params)
            QMessageBox.information(dialog, "완료", "'엑셀 가져옴' 프로필로 저장되었습니다.\n검색 패널에서 불러오세요.")
    except Exception as e:
        QMessageBox.critical(dialog, "오류", f"엑셀 읽기 실패: {e}")


def export_excel(dialog: Any) -> None:
    if not HAS_OPENPYXL:
        QMessageBox.critical(dialog, "오류", "openpyxl 라이브러리가 설치되지 않았습니다.\npip install openpyxl")
        return

    # Get results from MainWindow (parent)
    main_win = cast(Any, dialog.parent())
    if not main_win or not hasattr(main_win, 'all_results') or not main_win.all_results:
        QMessageBox.warning(dialog, "오류", "내보낼 검색 결과가 없습니다.")
        return

    fname, _ = QFileDialog.getSaveFileName(dialog, "엑셀로 저장", "flight_results.xlsx", "Excel Files (*.xlsx)")
    if not fname:
        return

    try:
        export_flights_to_excel(fname, main_win.all_results, sheet_title="검색결과")
        QMessageBox.information(dialog, "완료", "엑셀 파일로 저장되었습니다.")

    except Exception as e:
        QMessageBox.critical(dialog, "오류", f"엑셀 저장 실패: {e}")
