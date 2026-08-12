"""Shared flight-result export helpers."""

from __future__ import annotations

import csv
from typing import Any, Iterable


EXPORT_HEADERS = [
    "항공사",
    "오는편 항공사",
    "가격",
    "혜택가",
    "혜택 정보",
    "가는편 출발",
    "가는편 도착",
    "경유",
    "오는편 출발",
    "오는편 도착",
    "오는편 경유",
    "출처",
    "가는편 가격",
    "오는편 가격",
    # Extended fields (appended for backward-compatible column growth)
    "편명",
    "소요시간",
    "오는편 소요시간",
    "출발공항",
    "도착공항",
    "오는편 출발공항",
    "오는편 도착공항",
    "수하물",
    "오는편 수하물",
    "잔여석",
    "도착일+N",
    "오는편 도착일+N",
    "추천태그",
]
FORMULA_TRIGGER_PREFIXES = ("=", "+", "-", "@")


def sanitize_export_cell(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    stripped = value.lstrip()
    if stripped and stripped[0] in FORMULA_TRIGGER_PREFIXES:
        return "'" + value
    return value


def flight_export_headers() -> list[str]:
    return list(EXPORT_HEADERS)


def flight_to_export_row(flight: Any) -> list[Any]:
    return [
        getattr(flight, "airline", ""),
        getattr(flight, "return_airline", ""),
        getattr(flight, "price", 0),
        getattr(flight, "benefit_price", 0),
        getattr(flight, "benefit_label", ""),
        getattr(flight, "departure_time", ""),
        getattr(flight, "arrival_time", ""),
        getattr(flight, "stops", 0),
        getattr(flight, "return_departure_time", ""),
        getattr(flight, "return_arrival_time", ""),
        getattr(flight, "return_stops", 0),
        getattr(flight, "source", ""),
        getattr(flight, "outbound_price", 0),
        getattr(flight, "return_price", 0),
        getattr(flight, "flight_number", ""),
        getattr(flight, "duration", ""),
        getattr(flight, "return_duration", ""),
        getattr(flight, "departure_airport", ""),
        getattr(flight, "arrival_airport", ""),
        getattr(flight, "return_departure_airport", ""),
        getattr(flight, "return_arrival_airport", ""),
        getattr(flight, "baggage", ""),
        getattr(flight, "return_baggage", ""),
        getattr(flight, "seat_availability", 0),
        getattr(flight, "arrival_day_offset", 0),
        getattr(flight, "return_arrival_day_offset", 0),
        getattr(flight, "recommendation_tag", ""),
    ]


def flight_export_rows(flights: Iterable[Any]) -> list[list[Any]]:
    return [
        [sanitize_export_cell(value) for value in flight_to_export_row(flight)]
        for flight in flights
    ]


def export_flights_to_csv(filepath: str, flights: Iterable[Any]) -> None:
    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(flight_export_headers())
        writer.writerows(flight_export_rows(flights))


def export_flights_to_excel(filepath: str, flights: Iterable[Any], *, sheet_title: str = "검색 결과") -> None:
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet(title=sheet_title)
    else:
        ws.title = sheet_title

    headers = flight_export_headers()
    ws.append(headers)
    for row in flight_export_rows(flights):
        ws.append(row)

    for index in range(1, len(headers) + 1):
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = min(28, max(10, len(str(headers[index - 1])) + 2))

    wb.save(filepath)
